from pathlib import Path

def extract_text_from_file(file_path: str) -> str:
    """Extracts raw text from PDF, DOCX, or TXT file."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return _extract_from_pdf(path)
    elif extension in [".docx", ".doc"]:
        return _extract_from_docx(path)
    elif extension == ".txt":
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file format: {extension}")

def _extract_from_pdf(path: Path) -> str:
    import pypdf
    text = []
    with open(path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return "\n\n".join(text).strip()

def _extract_from_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(path)
        full_text = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text).strip()
    except Exception as e:
        # Fallback basic reading if docx library has issues with older .doc format
        with open(path, "rb") as f:
            raw = f.read()
            return raw.decode("utf-8", errors="ignore").strip()
